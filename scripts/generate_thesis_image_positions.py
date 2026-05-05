from __future__ import annotations

from pathlib import Path

from docx import Document


DOCX_PATH = Path("output/Graduation-thesis.docx")
OUT_PATH = Path("output/毕业论文插图位置清单.md")


def is_section_heading(text: str) -> bool:
    text = text.strip()
    if not text:
        return False
    if text.startswith("第 ") and "章" in text:
        return True
    marker = text.split(maxsplit=1)[0]
    return marker.count(".") >= 1 and all(part.isdigit() for part in marker.split("."))


def is_chapter_heading(text: str) -> bool:
    return text.startswith("第 ") and "章" in text


def is_figure_placeholder(text: str) -> bool:
    return text.startswith("图 ") and len(text) >= 5


def is_table_placeholder(text: str) -> bool:
    return text.startswith("表 ") and len(text) >= 5


def find_context(paragraphs: list[str], index: int) -> str:
    candidates: list[str] = []
    for cursor in range(index - 1, max(-1, index - 8), -1):
        text = paragraphs[cursor].strip()
        if text and not is_section_heading(text) and not is_figure_placeholder(text) and not is_table_placeholder(text):
            candidates.append(text)
        if len(candidates) >= 1:
            break
    if not candidates:
        return ""
    text = candidates[0]
    return text[:160] + ("……" if len(text) > 160 else "")


def main() -> None:
    doc = Document(DOCX_PATH)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    current_chapter = ""
    current_section = ""
    found_figures: list[dict[str, str]] = []
    found_tables: list[dict[str, str]] = []

    for idx, text in enumerate(paragraphs, start=1):
        if is_chapter_heading(text):
            current_chapter = text
            current_section = text
            continue
        if is_section_heading(text):
            current_section = text
            continue
        if is_figure_placeholder(text):
            found_figures.append(
                {
                    "paragraph": str(idx),
                    "chapter": current_chapter,
                    "section": current_section,
                    "caption": text,
                    "context": find_context(paragraphs, idx - 1),
                }
            )
        if is_table_placeholder(text):
            found_tables.append(
                {
                    "paragraph": str(idx),
                    "chapter": current_chapter,
                    "section": current_section,
                    "caption": text,
                    "context": find_context(paragraphs, idx - 1),
                }
            )

    # Based on the current thesis structure, these are the diagram positions that
    # should be kept even if the text is later reformatted. They correspond to
    # mechanism, architecture, and process descriptions rather than decorative images.
    recommended = [
        {
            "priority": "必插",
            "section": "3.1.1 链下日志存储与链上哈希存证模型",
            "caption": "图 3-1 链上链下混合存储模型",
            "content": "展示 logs/log_hash_records 等链下数据与 LogRegistry 链上哈希存证之间的关系。",
            "position": "放在本节说明“链下保存日志原文、链上保存哈希摘要”的段落之后。",
        },
        {
            "priority": "必插",
            "section": "3.1.2 日志采集、存证、审计、告警闭环流程",
            "caption": "图 3-2 可信日志审计闭环流程",
            "content": "展示“日志文件 -> Agent -> Server -> SQLite -> SHA-256 -> LogRegistry -> 审计比对 -> 告警 -> Web 展示”的闭环。",
            "position": "放在本节描述系统完整链路的段落之后。",
        },
        {
            "priority": "必插",
            "section": "3.4.1 数据库哈希、链上哈希与重算哈希三方比对",
            "caption": "图 3-3 三方哈希比对流程",
            "content": "展示 expectedHash、actualHash、onChainHash 的来源，以及 passed/failed/pending 的判断分支。",
            "position": "放在本节说明三方一致性判断逻辑之后。",
        },
        {
            "priority": "必插",
            "section": "4.4 系统总体架构设计",
            "caption": "图 4-1 系统总体架构图",
            "content": "展示 Agent、Server、SQLite、LogRegistry、Web 五类核心部分及数据流向。",
            "position": "放在本节介绍五类核心部分之后。",
        },
        {
            "priority": "必插",
            "section": "4.6 系统数据流与业务流程设计",
            "caption": "图 4-2 系统数据流图",
            "content": "分两条主流程展示日志采集与存证流程、审计与告警流程。",
            "position": "放在本节说明两条主流程之后。",
        },
        {
            "priority": "建议插",
            "section": "5.4 前端可视化模块实现",
            "caption": "图 5-1 系统前端总览页面",
            "content": "展示总览统计、日志数量、审计结果、告警数量、趋势图或分布图等页面效果。",
            "position": "放在 5.4 或 5.4.1 介绍前端总览与日志中心后。",
        },
        {
            "priority": "建议插",
            "section": "5.4.2 审计管理与告警管理实现",
            "caption": "图 5-2 审计管理与告警管理页面",
            "content": "展示审计记录列表、passed/failed/pending 状态、hash_mismatch 告警及处理状态。",
            "position": "放在 5.4.2 介绍审计结果和告警信息展示后。",
        },
        {
            "priority": "可选",
            "section": "5.6 性能实验与篡改检测实验",
            "caption": "图 5-3 批量审计耗时对比图",
            "content": "用柱状图展示 100/500/1000 条数据的平均审计耗时 3067.77/15659.44/35032.13 ms。",
            "position": "如果指导老师要求实验结果可视化，可放在 5.6.2 实验结果段落之后。",
        },
    ]

    lines: list[str] = []
    lines.append("# 毕业论文插图位置清单")
    lines.append("")
    lines.append(f"- 论文文件：`{DOCX_PATH.as_posix()}`")
    lines.append("- 说明：本清单只标注建议插入图片的位置，不直接修改论文正文。")
    lines.append("- 优先级建议：第 3 章机制图和第 4 章架构/流程图应优先补齐；第 5 章页面截图和实验图可根据页数与学校要求选择。")
    lines.append("")

    lines.append("## 一、论文中已有的图片占位")
    lines.append("")
    if found_figures:
        lines.append("| 序号 | 段落序号 | 所属章节 | 图片标题 | 建议插入位置依据 |")
        lines.append("|---:|---:|---|---|---|")
        for n, item in enumerate(found_figures, start=1):
            lines.append(
                f"| {n} | {item['paragraph']} | {item['section']} | {item['caption']} | {item['context']} |"
            )
    else:
        lines.append("当前 Word 正文中未扫描到以“图 X-X”开头的图片占位。")
    lines.append("")

    lines.append("## 二、建议补齐的插图位置")
    lines.append("")
    lines.append("| 序号 | 优先级 | 对应章节位置 | 建议图片标题 | 图片内容 | 放置位置 |")
    lines.append("|---:|---|---|---|---|---|")
    for n, item in enumerate(recommended, start=1):
        lines.append(
            f"| {n} | {item['priority']} | {item['section']} | {item['caption']} | {item['content']} | {item['position']} |"
        )
    lines.append("")

    lines.append("## 三、已有表格占位提示")
    lines.append("")
    lines.append("以下内容是表格，不属于图片；补论文图片时可以暂不处理。")
    lines.append("")
    if found_tables:
        lines.append("| 序号 | 段落序号 | 所属章节 | 表格标题 |")
        lines.append("|---:|---:|---|---|")
        for n, item in enumerate(found_tables, start=1):
            lines.append(f"| {n} | {item['paragraph']} | {item['section']} | {item['caption']} |")
    else:
        lines.append("当前正文中未扫描到以“表 X-X”开头的表格占位。")
    lines.append("")

    lines.append("## 四、插图补齐顺序建议")
    lines.append("")
    lines.append("1. 先补 `图 3-1`、`图 3-2`、`图 3-3`，这三张图对应论文核心机制。")
    lines.append("2. 再补 `图 4-1`、`图 4-2`，用于说明系统总体架构和业务流程。")
    lines.append("3. 最后根据篇幅补 `图 5-1`、`图 5-2` 页面截图；若实验部分需要更直观，可补 `图 5-3`。")
    lines.append("")

    OUT_PATH.write_text("\n".join(lines), encoding="utf-8")
    print(OUT_PATH.as_posix())
    print(f"figures={len(found_figures)}")
    print(f"table_placeholders={len(found_tables)}")


if __name__ == "__main__":
    main()
